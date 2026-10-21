"""pytest -q tests/test_api.py -- all tests use an isolated temporary SQLite DB."""
import io
import json
import sys
import time
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
import main
import store
from importer import extract

PASSWORD='Test-password-12345'
HEAD={'X-Konomi-Admin':'1'}

@pytest.fixture()
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA',tmp_path)
    monkeypatch.setenv('ADMIN_PASSWORD',PASSWORD)
    monkeypatch.delenv('GEMINI_API_KEY',raising=False)
    main.LIMITS.clear()
    main.WEATHER.update(ts=time.time(),data={'city':'Łódź','temp':8,'condition':'rain','source':'test','stale':False})
    with TestClient(main.app) as c: yield c

def auth(c):
    assert c.post('/api/admin/login',json={'username':'admin','password':PASSWORD},headers=HEAD).status_code==200

def item(c,id):return next(d for d in c.get('/api/admin/menu').json()['dishes'] if d['id']==id)

def update(c,id,**kw):
    body=item(c,id);body.pop('source',None);body.update(kw)
    return c.put('/api/admin/menu/'+id,json=body,headers=HEAD)

def preview(c,data,name='menu.csv',mapping=None):
    return c.post('/api/admin/import/preview',files={'file':(name,data)},data={'mapping':json.dumps(mapping or {})},headers=HEAD)

def test_public_and_admin_boundaries(client):
    assert len(client.get('/api/menu').json()['dishes'])==10
    assert client.get('/api/admin/menu').status_code==401
    assert client.put('/api/admin/settings',json={},headers=HEAD).status_code==401
    assert client.get('/api/health').json()['status']=='ok'

def test_auth_csrf_cookie_logout(client):
    assert client.post('/api/admin/login',json={'username':'admin','password':PASSWORD}).status_code==403
    assert client.post('/api/admin/login',json={'username':'admin','password':PASSWORD},headers={**HEAD,'Origin':'https://evil.test'}).status_code==403
    auth(client)
    assert client.get('/api/admin/me').status_code==200
    assert client.post('/api/admin/logout',headers=HEAD).status_code==200
    assert client.get('/api/admin/me').status_code==401

def test_key_is_encrypted_and_never_returned(client):
    auth(client);key='sample-not-real-key-123456'
    r=client.put('/api/admin/settings',json={'venue':'Test','model':'gemini-2.5-flash','discount':12,'gemini_key':key},headers=HEAD)
    assert r.status_code==200 and key not in r.text and r.json()['key_configured']
    with store.db() as c:value=c.execute("SELECT value FROM settings WHERE key='gemini_key'").fetchone()['value']
    assert key not in value and store.api_key()==key
    assert key not in client.get('/api/menu').text
    client.put('/api/admin/settings',json={'venue':'Test','model':'gemini-2.5-flash','discount':15},headers=HEAD)
    assert store.api_key()==key
    client.put('/api/admin/settings',json={'venue':'Test','model':'gemini-2.5-flash','clear_key':True},headers=HEAD)
    assert not store.api_key()

def test_invalid_money_and_drafts_and_version_conflict(client):
    auth(client)
    assert client.post('/api/admin/menu',json={'name':'Test','price':1.123},headers=HEAD).status_code==422
    r=client.post('/api/admin/menu',json={'name':'New soup','price':10,'category':'soup'},headers=HEAD)
    assert r.status_code==200
    assert not any(d['name']=='New soup' for d in client.get('/api/menu').json()['dishes'])
    d=item(client,'s1');d.pop('source')
    assert update(client,'s1',price=35).status_code==200
    assert client.put('/api/admin/menu/s1',json=d,headers=HEAD).status_code==409

def test_stop_list_blocks_checkout(client):
    auth(client);assert update(client,'s1',available=False).status_code==200
    assert client.post('/api/order/quote',json={'lines':[{'ids':['s1'],'qty':1}]}).status_code==409
    res=client.post('/api/chat',json={'message':'Рекомендации'}).json()
    assert 's1' not in res['dishes']

def test_discount_applies_to_combo_only(client):
    body={'lines':[{'ids':['s2','m1','b1'],'qty':1,'combo':True},{'ids':['d1'],'qty':1}]}
    r=client.post('/api/order/quote',json=body)
    assert r.status_code==200 and r.json()['total']==135.25
    assert r.json()['saving']==18.75
    assert client.post('/api/order/quote',json={'lines':[{'ids':['m1','m1','b1'],'qty':1,'combo':True}]}).status_code==422

def test_order_idempotency_and_stable_code(client):
    body={'lines':[{'ids':['s1'],'qty':1}],'idempotency_key':'my-test-check-123'}
    r=client.post('/api/orders',json=body);again=client.post('/api/orders',json=body)
    assert r.status_code==200 and r.json()==again.json()
    assert r.json()['status']=='created'

def test_unverified_never_means_allergy_safe(client):
    r=client.post('/api/chat',json={'message':'Не ем свинину и лактозу'}).json()
    assert r['dishes'] and 'pork' in r['excludes'] and 'lactose' in r['excludes']
    assert all(not ({'pork','lactose'} & set(d['allergens'])) for d in store.dishes() if d['id'] in r['dishes'])
    assert any(not r['dish_annotations'][i]['allergens_verified'] for i in r['dishes'])
    assert client.post('/api/order/quote',json={'lines':[{'ids':['m3'],'qty':1}],'excludes':['pork']}).status_code==200
    assert client.post('/api/order/quote',json={'lines':[{'ids':['m1'],'qty':1}],'excludes':['pork']}).status_code==409
    auth(client);update(client,'s1',allergens_verified=True);update(client,'m3',allergens_verified=True)
    r=client.post('/api/chat',json={'message':'Не ем свинину и лактозу'}).json()
    assert {'s1','m3'}.issubset(r['dishes'])
    assert set(r['dishes'])=={d['id'] for d in store.dishes() if not set(d['allergens']) & {'pork','lactose'}}

def test_exclusion_persistence_and_mentions(client):
    assert not main.exclusions('Хочу свинину')
    assert 'gluten' in main.exclusions('gluten-free please')
    assert 'pork' in main.exclusions('no pork')
    r=client.post('/api/chat',json={'message':'Собери комбо','history':[{'role':'user','content':'Без свинины'}]}).json()
    assert r['combo'] and r['excludes']==['pork']
    assert all('pork' not in d['allergens'] for d in store.dishes() if d['id'] in r['combo']['ids'])

def test_csv_mapping_raw_details_draft_commit(client):
    auth(client)
    raw='Название;Цена;Вес;Состав;Регион\nТестовый суп;31,50;350g;Капуста;Татры\n'.encode('utf-8-sig')
    r=preview(client,raw);assert r.status_code==200
    p=r.json();assert p['rows'][0]['dish']['price']==31.5
    assert p['rows'][0]['raw']['Регион']=='Татры'
    res=client.post('/api/admin/import/commit',json={'token':p['token']},headers=HEAD)
    assert res.status_code==200 and res.json()['added']==1
    d=next(d for d in client.get('/api/admin/menu').json()['dishes'] if d['name']=='Тестовый суп')
    assert d['status']=='draft' and not d['allergens_verified']
    assert d['source']['raw']['Регион']=='Татры' and d['source']['filename']=='menu.csv'
    assert client.post('/api/admin/import/commit',json={'token':p['token']},headers=HEAD).status_code==409

def test_preview_errors_block_import_and_custom_map(client):
    auth(client);raw=b'Product,Amount,Cost\nDish,350,NaN\n'
    r=preview(client,raw,mapping={'name':'Product','price':'Cost','weight':'Amount'})
    assert r.status_code==200 and r.json()['rows'][0]['errors']
    assert client.post('/api/admin/import/commit',json={'token':r.json()['token']},headers=HEAD).status_code==422

def test_duplicate_import_safe_update(client):
    auth(client)
    raw=json.dumps([{'id':'s1','name':'Updated soup','price':45,'category':'soup'}]).encode()
    p=preview(client,raw,'menu.json').json()
    assert client.post('/api/admin/import/commit',json={'token':p['token']},headers=HEAD).json()['skipped']==1
    p=preview(client,raw,'menu.json').json()
    assert client.post('/api/admin/import/commit',json={'token':p['token'],'strategy':'update'},headers=HEAD).json()['updated']==1
    assert item(client,'s1')['status']=='draft'

def test_xlsx_and_docx_parsing(client):
    auth(client)
    from openpyxl import Workbook
    w=Workbook();w.active.append(['name','price']);w.active.append(['Test XLSX',27.55]);f=io.BytesIO();w.save(f)
    r=preview(client,f.getvalue(),'menu.xlsx');assert r.status_code==200 and r.json()['rows'][0]['dish']['price']==27.55
    from docx import Document
    doc=Document();doc.add_paragraph('Rosół 31 PLN 350g');f=io.BytesIO();doc.save(f)
    r=preview(client,f.getvalue(),'menu.docx');assert r.status_code==200 and r.json()['needs_ai']
    assert 'Rosół' in r.json()['text_preview']

def test_import_file_limits_and_types(client):
    auth(client)
    assert preview(client,b'x','menu.exe').status_code==415
    assert preview(client,b'x'*5_000_001).status_code==413
    assert preview(client,b'{broken','menu.json').status_code==422

def test_ai_selects_full_catalog_then_ids_are_validated(client,monkeypatch):
    async def fake(request,prompt):return {'reply':'Try these dishes','dishes':['not-real','s1'],'combo':{'ids':['s1','m1','m2']}}
    monkeypatch.setattr(main,'generate',fake);monkeypatch.setattr(store,'api_key',lambda provider='gemini':'test' if provider=='gemini' else '')
    r=client.post('/api/chat',json={'message':'Hi'}).json()
    assert r['dishes']==['s1'] and 'not-real' not in r['dishes'] and r['combo']['discount']==15 and r['engine']=='gemini'

def test_weather_does_not_invent_temperature(client,monkeypatch):
    main.WEATHER.update(ts=0,data=None)
    async def fail(*args,**kwargs):raise main.httpx.ConnectError('offline')
    monkeypatch.setattr(client.app.state.http,'get',fail)
    r=client.get('/api/weather').json()
    assert r['temp'] is None and r['source']=='not-configured'

def test_password_rotation_invalidates_all_sessions(client):
    auth(client)
    assert client.post('/api/admin/password',json={'current':PASSWORD,'password':'New-password-9999'},headers=HEAD).status_code==200
    assert client.get('/api/admin/me').status_code==401
    assert client.post('/api/admin/login',json={'username':'admin','password':PASSWORD},headers=HEAD).status_code==401

def test_request_validation(client):
    assert client.post('/api/chat',json={'message':'x'*2001}).status_code==422
    assert client.post('/api/order/quote',json={'lines':[{'ids':['s1'],'qty':-1}]}).status_code==422
