"""Menu import is deliberately a staged operation: parse, map, review, commit.
Unrecognized source columns are preserved in source.raw; no allergen is invented.
"""
import csv
import io
import json
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from fastapi import HTTPException

FIELDS = ['id','name','category','price','weight','allergens','desc','image','ingredients','details']
ALIASES = {
 'id':['id','sku','артикул','код'], 'name':['name','название','nazwa','dish','блюдо'],
 'category':['category','категория','kategoria','cat'], 'price':['price','цена','cena','price pln','цена pln'],
 'weight':['weight','вес','граммовка','waga','объем'], 'allergens':['allergens','аллергены','alergeny'],
 'desc':['desc','description','описание','opis'], 'image':['image','photo','фото','zdjecie','url'],
 'ingredients':['ingredients','состав','sklad'], 'details':['details','детали','uwagi']
}

def normalized(text):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(text).lower().strip()) if not unicodedata.combining(c))

def extract(filename, data):
    ext = Path(filename).suffix.lower()
    if ext in ('.xlsx','.docx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > 40_000_000:
                    raise HTTPException(413, 'Распакованный файл слишком большой')
        except zipfile.BadZipFile:
            raise HTTPException(422, 'Поврежденный Office-файл')
    if ext == '.json':
        raw = json.loads(data.decode('utf-8-sig'))
        raw = raw.get('dishes',raw.get('menu',raw.get('items',[]))) if isinstance(raw,dict) else raw
        if not isinstance(raw,list) or any(not isinstance(r,dict) for r in raw):
            raise HTTPException(422,'JSON должен содержать массив объектов блюд')
        return raw, None
    if ext == '.csv':
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('cp1251')
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        return list(csv.DictReader(io.StringIO(text), dialect=dialect)), None
    if ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        sheet = book.active
        if sheet.max_row and sheet.max_row > 1001:
            book.close(); raise HTTPException(422,'Максимум 1000 строк')
        rows = list(sheet.iter_rows(values_only=True))
        book.close()
        if not rows: return [], None
        headers = [str(x or f'column_{i+1}') for i,x in enumerate(rows[0])]
        return [dict(zip(headers,r)) for r in rows[1:] if any(x is not None for x in r)], None
    if ext == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(data))
        if len(pdf.pages)>100: raise HTTPException(422,'PDF: не больше 100 страниц')
        return [], '\n'.join((p.extract_text() or '') for p in pdf.pages)[:80000]
    if ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return [], '\n'.join(parts)[:80000]
    raise HTTPException(415,'Поддерживаются JSON, CSV, XLSX, PDF и DOCX')

def mapping_for(rows):
    headers = list(dict.fromkeys(k for r in rows for k in r))
    return {f: next((h for h in headers if normalized(h) in [normalized(a) for a in ALIASES[f]]), '') for f in FIELDS}

def money(value):
    if isinstance(value,bool): raise ValueError()
    text = str(value if value is not None else '').replace('\u00a0','').replace(' ','')
    text = re.sub(r'(?i)(PLN|zł|zl)$','',text).replace(',','.')
    try: n = Decimal(text)
    except InvalidOperation: raise ValueError('Неверная цена')
    if not n.is_finite() or not 0 <= n <= 100000 or n.as_tuple().exponent < -2:
        raise ValueError('Цена должна быть от 0 до 100000 с точностью до гроша')
    return float(n)

def make_preview(rows, mapping):
    if len(rows)>1000: raise HTTPException(422,'Максимум 1000 блюд за импорт')
    result=[]
    for i, raw in enumerate(rows):
        get=lambda field: raw.get(mapping.get(field,''))
        errors=[]; warnings=[]
        try: price=money(get('price'))
        except ValueError: price=0; errors.append('Укажите корректную цену в PLN')
        name=str(get('name') or '').strip()
        if not name: errors.append('Нет названия')
        cat=normalized(get('category') or '')
        cats={'soup':'soup','супы':'soup','суп':'soup','zupy':'soup','main':'main','горячее':'main','основные':'main','dania glowne':'main','beverage':'beverage','напитки':'beverage','napoje':'beverage','dessert':'dessert','десерты':'dessert','desery':'dessert','starter':'starter','закуски':'starter'}
        category=cats.get(cat,'other')
        if category=='other': warnings.append('Проверьте категорию')
        desc=get('desc') or ''
        if not isinstance(desc,dict): desc={'ru':str(desc)}
        al=get('allergens') or []
        if not isinstance(al,list): al=[a.strip().lower() for a in re.split('[,;|]',str(al)) if a.strip()]
        a_map={'глютен':'gluten','лактоза':'lactose','свинина':'pork','рыба':'fish','молоко':'milk','яйца':'egg','орехи':'nuts','gluten':'gluten','laktoza':'lactose','wieprzowina':'pork'}
        al=[a_map.get(str(a).lower(),str(a).lower()) for a in al]
        warnings.append('Состав и аллергены требуют подтверждения рестораном')
        image=str(get('image') or '')
        if image and not image.startswith(('https://','http://')):
            image=''; warnings.append('Фото: нужен http(s) URL')
        dish={'id':str(get('id') or ''),'name':name,'category':category,'price':price,'weight':str(get('weight') or ''),'allergens':al,'desc':desc,'image':image,'ingredients':str(get('ingredients') or ''),'details':str(get('details') or ''),'status':'draft','available':True,'allergens_verified':False,'tags':[]}
        result.append({'row':i+1,'dish':dish,'errors':errors,'warnings':warnings,'raw':raw})
    return result
